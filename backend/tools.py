"""Implementation of every PDF tool. Each tool: fn(files: list[Upload], opts: dict, workdir: str) -> Result"""
import base64
import concurrent.futures as cf
import difflib
import io
import json
import math
import os
import re
import secrets
import urllib.parse
import urllib.request
from collections import Counter
from pathlib import Path

import pymupdf
from PIL import Image, ImageOps, ImageEnhance, ImageFilter

from .core import (UserError, Result, Upload, open_pdf, pdf_bytes, decrypted_path, require_pdfs,
                   parse_pages, parse_ranges, hex_to_rgb, rgb_to_hex, insert_text_box, make_text_stamp,
                   normalize_page, libreoffice_convert, gs_bin, chrome_bin, tesseract_bin, run, safe_name)

try:
    import pillow_heif  # optional HEIC support
    pillow_heif.register_heif_opener()
except Exception:
    pass

IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tif", ".tiff", ".webp", ".heic", ".heif", ".jfif"}


def _b(v, default=False):
    if isinstance(v, bool):
        return v
    if v is None:
        return default
    return str(v).lower() in ("1", "true", "yes", "on")


def _num(v, default, lo=None, hi=None, cast=float):
    try:
        x = cast(v)
        if isinstance(x, float) and (math.isnan(x) or math.isinf(x)):
            raise ValueError
    except (TypeError, ValueError):
        x = default
    if lo is not None:
        x = max(lo, x)
    if hi is not None:
        x = min(hi, x)
    return x


def _single_or_zip(files, info=None):
    return Result(files=files, info=info or {})


# =========================================================================== #
# ORGANIZE
# =========================================================================== #
def merge(files, opts, wd):
    pdfs = require_pdfs(files, 2)
    out = pymupdf.open()
    toc = []
    for up in pdfs:
        doc = open_pdf(up)
        start = out.page_count
        out.insert_pdf(doc)
        toc.append([1, Path(up.name).stem, start + 1])
        doc.close()
    if _b(opts.get("bookmarks"), True):
        out.set_toc(toc)
    return Result([(f"{pdfs[0].stem}_merged.pdf", pdf_bytes(out))], {"pages": out.page_count})


def split(files, opts, wd):
    pdfs = require_pdfs(files, 1)
    mode = opts.get("mode", "ranges")
    results = []
    for up in pdfs:
        doc = open_pdf(up)
        n = doc.page_count
        if mode == "ranges":
            groups = parse_ranges(opts.get("ranges") or f"1-{n}", n)
        elif mode == "fixed":
            k = _num(opts.get("every"), 1, 1, n, int)
            groups = [list(range(i, min(i + k, n))) for i in range(0, n, k)]
        elif mode == "all":
            groups = [[i] for i in range(n)]
        elif mode == "extract":
            pages = parse_pages(opts.get("pages"), n)
            groups = [pages] if _b(opts.get("merge_extract"), True) else [[p] for p in pages]
        else:
            raise UserError("Unknown split mode.")
        if mode == "ranges" and _b(opts.get("merge")):
            flat = [p for g in groups for p in g]
            groups = [flat]
        for g in groups:
            part = pymupdf.open()
            for p in g:
                part.insert_pdf(doc, from_page=p, to_page=p)
            label = f"{g[0] + 1}" if len(g) == 1 else f"{g[0] + 1}-{g[-1] + 1}"
            results.append((f"{up.stem}_{label}.pdf", pdf_bytes(part)))
            part.close()
        doc.close()
    return Result(results, {"files": len(results)})


def remove_pages(files, opts, wd):
    up = require_pdfs(files, 1, 1)[0]
    doc = open_pdf(up)
    pages = parse_pages(opts.get("pages"), doc.page_count)
    if len(pages) >= doc.page_count:
        raise UserError("You cannot remove all pages of the document.")
    doc.delete_pages(pages)
    return Result([(f"{up.stem}_removed.pdf", pdf_bytes(doc))], {"pages": doc.page_count})


def extract_pages(files, opts, wd):
    up = require_pdfs(files, 1, 1)[0]
    doc = open_pdf(up)
    pages = parse_pages(opts.get("pages"), doc.page_count)
    if _b(opts.get("separate")):
        out = []
        for p in pages:
            part = pymupdf.open()
            part.insert_pdf(doc, from_page=p, to_page=p)
            out.append((f"{up.stem}_page{p + 1}.pdf", pdf_bytes(part)))
        return Result(out)
    new = pymupdf.open()
    for p in pages:
        new.insert_pdf(doc, from_page=p, to_page=p)
    return Result([(f"{up.stem}_extracted.pdf", pdf_bytes(new))], {"pages": new.page_count})


def organize(files, opts, wd):
    """opts.pages = [{file: idx, page: idx, rotate: deg} | {blank: true, width, height}]"""
    pdfs = require_pdfs(files, 1)
    layout = opts.get("pages")
    if not isinstance(layout, list) or not layout:
        raise UserError("The document must contain at least one page.")
    docs = [open_pdf(up) for up in pdfs]
    out = pymupdf.open()
    for item in layout:
        if item.get("blank"):
            ref = out[-1].rect if out.page_count else docs[0][0].rect
            out.new_page(width=_num(item.get("width"), ref.width, 10, 14400),
                         height=_num(item.get("height"), ref.height, 10, 14400))
        else:
            fi = _num(item.get("file"), 0, 0, len(docs) - 1, int)
            pi = int(item.get("page", 0))
            if pi < 0 or pi >= docs[fi].page_count:
                raise UserError("Invalid page reference.")
            out.insert_pdf(docs[fi], from_page=pi, to_page=pi)
        rot = int(item.get("rotate", 0)) % 360
        if rot:
            pg = out[-1]
            pg.set_rotation((pg.rotation + rot) % 360)
    return Result([(f"{pdfs[0].stem}_organized.pdf", pdf_bytes(out))], {"pages": out.page_count})


def rotate(files, opts, wd):
    pdfs = require_pdfs(files, 1)
    angle = int(_num(opts.get("angle"), 90, cast=int)) % 360
    if angle % 90:
        raise UserError("Rotation must be a multiple of 90 degrees.")
    per_file = opts.get("per_file") or {}
    out = []
    for i, up in enumerate(pdfs):
        doc = open_pdf(up)
        a = int(per_file.get(str(i), angle)) % 360 if per_file else angle
        for p in parse_pages(opts.get("pages"), doc.page_count):
            pg = doc[p]
            pg.set_rotation((pg.rotation + a) % 360)
        out.append((f"{up.stem}_rotated.pdf", pdf_bytes(doc)))
        doc.close()
    return Result(out)


# =========================================================================== #
# OPTIMIZE
# =========================================================================== #
COMPRESS = {
    "low": dict(threshold=200, target=150, quality=80, gs="/printer"),
    "recommended": dict(threshold=150, target=110, quality=65, gs="/ebook"),
    "extreme": dict(threshold=96, target=72, quality=40, gs="/screen"),
}


def compress(files, opts, wd):
    pdfs = require_pdfs(files, 1)
    level = COMPRESS.get(opts.get("level", "recommended"), COMPRESS["recommended"])
    grayscale = _b(opts.get("grayscale"))
    out, before_total, after_total = [], 0, 0
    for up in pdfs:
        original = Path(up.path).read_bytes()
        candidates = []
        doc = open_pdf(up)
        try:
            doc.rewrite_images(dpi_threshold=level["threshold"], dpi_target=level["target"],
                               quality=level["quality"], lossy=True, lossless=True, bitonal=False,
                               set_to_gray=grayscale)
        except Exception:
            pass
        try:
            doc.subset_fonts()
        except Exception:
            pass
        doc.scrub(metadata=False, xml_metadata=False, thumbnails=True, reset_fields=False,
                  reset_responses=False, remove_links=False, attached_files=False, clean_pages=False,
                  embedded_files=False, hidden_text=False, javascript=False, redactions=False)
        candidates.append(doc.tobytes(garbage=4, deflate=True, deflate_images=True, deflate_fonts=True,
                                      clean=True, use_objstms=1, encryption=pymupdf.PDF_ENCRYPT_NONE))
        doc.close()
        gs = gs_bin()
        if gs and opts.get("level") in ("recommended", "extreme") or (gs and grayscale):
            src = decrypted_path(up, wd)
            dst = os.path.join(wd, f"gs_{secrets.token_hex(4)}.pdf")
            cmd = [gs, "-sDEVICE=pdfwrite", "-dCompatibilityLevel=1.5", f"-dPDFSETTINGS={level['gs']}",
                   "-dNOPAUSE", "-dQUIET", "-dBATCH", "-dDetectDuplicateImages=true"]
            if grayscale:
                cmd += ["-sColorConversionStrategy=Gray", "-dProcessColorModel=/DeviceGray"]
            cmd += [f"-sOutputFile={dst}", src]
            try:
                run(cmd, timeout=600)
                if os.path.exists(dst) and os.path.getsize(dst) > 0:
                    chk = pymupdf.open(dst)
                    if chk.page_count == pymupdf.open(src).page_count:
                        candidates.append(Path(dst).read_bytes())
                    chk.close()
            except Exception:
                pass
        best = min(candidates, key=len)
        if len(best) >= len(original) and not grayscale:
            best = original if not up.password else candidates[0]
        before_total += len(original)
        after_total += len(best)
        out.append((f"{up.stem}_compressed.pdf", best))
    saved = max(0, before_total - after_total)
    return Result(out, {"before": before_total, "after": after_total,
                        "saved_percent": round(saved * 100 / before_total, 1) if before_total else 0})


def repair(files, opts, wd):
    pdfs = require_pdfs(files, 1)
    out = []
    for up in pdfs:
        data, method = None, None
        try:
            doc = pymupdf.open(up.path, filetype="pdf")
            if doc.needs_pass:
                if not up.password or not doc.authenticate(up.password):
                    raise UserError(f"“{up.name}” is password protected. Please enter its password.")
            if doc.page_count > 0:
                data = doc.tobytes(garbage=4, deflate=True, clean=True, encryption=pymupdf.PDF_ENCRYPT_NONE)
                method = "rebuilt"
            doc.close()
        except UserError:
            raise
        except Exception:
            data = None
        if data is None:
            try:
                from pypdf import PdfReader, PdfWriter
                r = PdfReader(up.path, strict=False)
                if r.is_encrypted and up.password:
                    r.decrypt(up.password)
                w = PdfWriter()
                for p in r.pages:
                    w.add_page(p)
                buf = io.BytesIO()
                w.write(buf)
                if len(r.pages):
                    data, method = buf.getvalue(), "recovered"
            except Exception:
                data = None
        if data is None and gs_bin():
            dst = os.path.join(wd, f"rep_{secrets.token_hex(4)}.pdf")
            run([gs_bin(), "-o", dst, "-sDEVICE=pdfwrite", "-dPDFSTOPONERROR=false", up.path], timeout=600)
            if os.path.exists(dst) and os.path.getsize(dst) > 200:
                data, method = Path(dst).read_bytes(), "reconstructed"
        if data is None:
            raise UserError(f"Sorry, “{up.name}” is too damaged to be repaired.")
        out.append((f"{up.stem}_repaired.pdf", data))
    return Result(out)


def _tess_langs():
    exe = tesseract_bin()
    if not exe:
        return []
    proc = run([exe, "--list-langs"], timeout=30)
    text = (proc.stdout or b"").decode() + (proc.stderr or b"").decode()
    langs = [l.strip() for l in text.splitlines()[1:] if l.strip() and " " not in l.strip()]
    return sorted(set(langs) - {"osd"})


def ocr_languages():
    return _tess_langs()


def ocr(files, opts, wd):
    exe = tesseract_bin()
    if not exe:
        raise UserError("OCR needs the free Tesseract OCR engine. Install it (see README) and restart the app.")
    import pytesseract
    pytesseract.pytesseract.tesseract_cmd = exe
    pdfs = require_pdfs(files, 1)
    langs = opts.get("languages") or ["eng"]
    if isinstance(langs, str):
        langs = [langs]
    available = _tess_langs()
    missing = [l for l in langs if l not in available]
    if missing:
        raise UserError(f"OCR language(s) not installed: {', '.join(missing)}. Installed: {', '.join(available)}")
    lang = "+".join(langs)
    dpi = 300
    skip_text = _b(opts.get("skip_text"), True)
    out, total_pages = [], 0

    def ocr_page(png):
        img = Image.open(io.BytesIO(png))
        return pytesseract.image_to_pdf_or_hocr(img, extension="pdf", lang=lang,
                                                config=f"--dpi {dpi} -c textonly_pdf=1")

    for up in pdfs:
        doc = open_pdf(up)
        jobs = {}
        for page in doc:
            if skip_text and len(page.get_text().strip()) > 20:
                continue
            normalize_page(page)
            zoom = dpi / 72
            pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), colorspace=pymupdf.csGRAY)
            jobs[page.number] = pix.tobytes("png")
        with cf.ThreadPoolExecutor(max_workers=max(1, min(4, os.cpu_count() or 2))) as ex:
            futures = {ex.submit(ocr_page, png): pno for pno, png in jobs.items()}
            layers = {}
            for fut in cf.as_completed(futures):
                layers[futures[fut]] = fut.result()
        for pno, pdfdata in layers.items():
            layer = pymupdf.open("pdf", pdfdata)
            page = doc[pno]
            page.show_pdf_page(page.rect, layer, 0, overlay=True)
            layer.close()
        total_pages += len(jobs)
        out.append((f"{up.stem}_ocr.pdf", pdf_bytes(doc)))
        doc.close()
    return Result(out, {"ocr_pages": total_pages})


# =========================================================================== #
# CONVERT TO PDF
# =========================================================================== #
PAGE_SIZES = {"A4": (595.28, 841.89), "Letter": (612, 792), "Legal": (612, 1008), "A3": (841.89, 1190.55),
              "A5": (419.53, 595.28)}


def _prepare_image(path, filt="none"):
    try:
        img = Image.open(path)
        img = ImageOps.exif_transpose(img)
    except Exception:
        raise UserError(f"“{Path(path).name}” is not a supported image.")
    if getattr(img, "n_frames", 1) > 1:
        img.seek(0)
    if img.mode in ("RGBA", "LA", "P"):
        img = img.convert("RGBA")
        bg = Image.new("RGB", img.size, (255, 255, 255))
        bg.paste(img, mask=img.split()[-1])
        img = bg
    elif img.mode not in ("RGB", "L"):
        img = img.convert("RGB")
    if filt == "gray":
        img = ImageOps.autocontrast(img.convert("L"), cutoff=1)
    elif filt == "bw":
        g = ImageOps.autocontrast(img.convert("L"), cutoff=2).filter(ImageFilter.SHARPEN)
        img = g.point(lambda v: 255 if v > 150 else 0).convert("L")
    elif filt == "enhance":
        img = ImageEnhance.Contrast(ImageOps.autocontrast(img, cutoff=1)).enhance(1.3)
        img = ImageEnhance.Sharpness(img).enhance(1.5)
    buf = io.BytesIO()
    if img.mode == "L" and filt == "bw":
        img.save(buf, "PNG", optimize=True)
    else:
        img.save(buf, "JPEG", quality=92)
    return img.size, buf.getvalue()


def images_to_pdf(files, opts, wd):
    if not files:
        raise UserError("Please select at least one image.")
    bad = [f.name for f in files if f.ext not in IMAGE_EXTS]
    if bad:
        raise UserError("These files are not supported images: " + ", ".join(bad))
    orientation = opts.get("orientation", "portrait")
    size_key = opts.get("pagesize", "fit")
    margin = {"none": 0, "small": 20, "big": 50}.get(opts.get("margin", "none"), 0)
    filt = opts.get("filter", "none")
    merge_all = _b(opts.get("merge"), True)
    docs = []
    doc = pymupdf.open()
    for up in files:
        (w, h), data = _prepare_image(up.path, filt)
        if size_key == "fit":
            pw, ph = w * 72 / 96, h * 72 / 96
            scale = min(1.0, 14400 / max(pw, ph))
            pw, ph = pw * scale + 2 * margin, ph * scale + 2 * margin
        else:
            pw, ph = PAGE_SIZES.get(size_key, PAGE_SIZES["A4"])
            land = (orientation == "landscape") or (orientation == "auto" and w > h)
            if land:
                pw, ph = ph, pw
        page = doc.new_page(width=pw, height=ph)
        rect = pymupdf.Rect(margin, margin, pw - margin, ph - margin)
        page.insert_image(rect, stream=data, keep_proportion=True)
        if not merge_all:
            docs.append((f"{up.stem}.pdf", pdf_bytes(doc)))
            doc = pymupdf.open()
    if merge_all:
        name = "scan.pdf" if opts.get("scan") else f"{files[0].stem}.pdf" if len(files) == 1 else "images.pdf"
        return Result([(name, pdf_bytes(doc))], {"pages": doc.page_count})
    return Result(docs)


OFFICE_EXTS = {
    "word": {".doc", ".docx", ".odt", ".rtf", ".txt", ".wpd", ".docm", ".dot", ".dotx"},
    "powerpoint": {".ppt", ".pptx", ".odp", ".pps", ".ppsx", ".pptm", ".key"},
    "excel": {".xls", ".xlsx", ".ods", ".csv", ".xlsm", ".xlsb"},
}


def office_to_pdf(kind):
    def tool(files, opts, wd):
        if not files:
            raise UserError("Please select at least one file.")
        allowed = OFFICE_EXTS[kind]
        bad = [f.name for f in files if f.ext not in allowed]
        if bad:
            raise UserError(f"Unsupported file type: {', '.join(bad)}. Allowed: {', '.join(sorted(allowed))}")
        out = []
        for up in files:
            src_dir = os.path.join(wd, f"src_{secrets.token_hex(4)}")
            os.makedirs(src_dir)
            src = os.path.join(src_dir, f"{up.stem[:80] or 'file'}{up.ext}")
            os.replace(up.path, src)
            up.path = src
            res = libreoffice_convert(src, "pdf", src_dir)
            out.append((f"{up.stem}.pdf", Path(res).read_bytes()))
        return Result(out)
    return tool


def _run_chrome(cmd, out_path, timeout=60):
    """Run headless Chrome; some builds never exit after printing, so poll for the output file."""
    import subprocess, time as _t, platform as _p
    kw = {"creationflags": 0x08000000} if _p.system() == "Windows" else {}
    try:
        proc = subprocess.Popen(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, **kw)
    except OSError:
        return
    start, last_size, stable = _t.time(), -1, 0
    try:
        while _t.time() - start < timeout:
            if proc.poll() is not None:
                break
            if os.path.exists(out_path):
                size = os.path.getsize(out_path)
                stable = stable + 1 if size == last_size and size > 0 else 0
                last_size = size
                if stable >= 3:
                    break
            _t.sleep(0.3)
    finally:
        if proc.poll() is None:
            proc.kill()
            try:
                proc.wait(5)
            except Exception:
                pass
    if os.path.exists(out_path):
        try:
            pymupdf.open(out_path).close()
        except Exception:
            os.remove(out_path)


def html_to_pdf(files, opts, wd):
    url = (opts.get("url") or "").strip()
    target = None
    if files:
        up = files[0]
        if up.ext not in (".html", ".htm", ".xhtml", ".mhtml"):
            raise UserError("Please upload an .html file or enter a URL.")
        target = Path(up.path).resolve().as_uri()
        name = up.stem
    elif url:
        if not re.match(r"^https?://", url, re.I):
            url = "https://" + url
        parsed = urllib.parse.urlparse(url)
        if parsed.scheme not in ("http", "https") or not parsed.netloc:
            raise UserError("Please enter a valid web address, e.g. https://example.com")
        from .public_proxy import validate_public_url
        try:
            validate_public_url(url)
        except (ValueError, OSError) as e:
            raise UserError(str(e))
        target = url
        name = safe_name(parsed.netloc.replace("www.", "")) or "webpage"
    else:
        raise UserError("Please enter a URL or upload an HTML file.")
    out_path = os.path.join(wd, "page.pdf")
    from .public_proxy import proxy_url, PublicRedirectHandler
    conversion_proxy = proxy_url()
    exe = chrome_bin()
    if exe:
        profile = os.path.join(wd, "chrome_profile")
        cmd = [exe, "--headless=new", "--disable-gpu", "--no-sandbox", "--hide-scrollbars",
               "--disable-dev-shm-usage", f"--proxy-server={conversion_proxy}", "--proxy-bypass-list=<-loopback>",
               "--disable-extensions",
               f"--user-data-dir={profile}", f"--print-to-pdf={out_path}"]
        wait_ms = int(_num(opts.get("wait"), 3, 0, 30)) * 1000
        if wait_ms > 0:
            cmd.append(f"--virtual-time-budget={wait_ms}")
        if not _b(opts.get("header_footer")):
            cmd.append("--no-pdf-header-footer")
        cmd.append(target)
        _run_chrome(cmd, out_path, timeout=60 + int(_num(opts.get("wait"), 3, 0, 30)))
    if not os.path.exists(out_path) or os.path.getsize(out_path) == 0:
        # Fallback: LibreOffice
        if target.startswith("http"):
            try:
                req = urllib.request.Request(target, headers={"User-Agent": "Mozilla/5.0"})
                opener = urllib.request.build_opener(urllib.request.ProxyHandler({"http": conversion_proxy, "https": conversion_proxy}), PublicRedirectHandler())
                html_bytes = opener.open(req, timeout=30).read(20 * 1024 * 1024 + 1)
                if len(html_bytes) > 20 * 1024 * 1024:
                    raise UserError("The web page exceeds the 20 MB conversion limit.")
            except Exception as e:
                raise UserError(f"Could not open the web page: {e}")
            src = os.path.join(wd, "page.html")
            Path(src).write_bytes(html_bytes)
        else:
            src = files[0].path
            if not src.endswith(".html"):
                os.replace(src, src + ".html")
                src = src + ".html"
        res = libreoffice_convert(src, "pdf", wd)
        out_path = res
    data = Path(out_path).read_bytes()
    doc = pymupdf.open("pdf", data)
    if doc.page_count == 0:
        raise UserError("The page could not be converted.")
    return Result([(f"{name}.pdf", data)], {"pages": doc.page_count})


# =========================================================================== #
# CONVERT FROM PDF
# =========================================================================== #
def pdf_to_images(files, opts, wd):
    pdfs = require_pdfs(files, 1)
    mode = opts.get("mode", "pages")
    fmt = "png" if opts.get("format") == "png" else "jpg"
    dpi = int(_num(opts.get("dpi"), 150, 36, 600))
    out = []
    for up in pdfs:
        doc = open_pdf(up)
        if mode == "extract":
            seen = set()
            for page in doc:
                for img in page.get_images(full=True):
                    xref, smask = img[0], img[1]
                    if xref in seen:
                        continue
                    seen.add(xref)
                    try:
                        if smask:
                            base = pymupdf.Pixmap(doc, xref)
                            mask = pymupdf.Pixmap(doc, smask)
                            if base.alpha:
                                base = pymupdf.Pixmap(base, 0)
                            pix = pymupdf.Pixmap(base, mask)
                            if pix.width < 16 or pix.height < 16:
                                continue
                            out.append((f"{up.stem}_p{page.number + 1}_img{len(seen)}.png", pix.tobytes("png")))
                        else:
                            info = doc.extract_image(xref)
                            if not info or info["width"] < 16 or info["height"] < 16:
                                continue
                            ext = info["ext"]
                            data = info["image"]
                            if ext not in ("jpg", "jpeg", "png"):
                                pix = pymupdf.Pixmap(doc, xref)
                                if pix.n - pix.alpha >= 4:
                                    pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
                                data, ext = pix.tobytes("png"), "png"
                            out.append((f"{up.stem}_p{page.number + 1}_img{len(seen)}.{ext}", data))
                    except Exception:
                        continue
        else:
            pages = parse_pages(opts.get("pages"), doc.page_count)
            zoom = dpi / 72
            for p in pages:
                pix = doc[p].get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False)
                if fmt == "png":
                    data = pix.tobytes("png")
                else:
                    data = pix.tobytes("jpg", jpg_quality=int(_num(opts.get("quality"), 90, 30, 100)))
                out.append((f"{up.stem}_page{p + 1}.{fmt}", data))
        doc.close()
    if not out:
        raise UserError("No images were found in this PDF." if mode == "extract" else "Nothing to convert.")
    return Result(out, {"images": len(out)})


def pdf_to_word(files, opts, wd):
    from pdf2docx import Converter
    pdfs = require_pdfs(files, 1)
    out = []
    for up in pdfs:
        src = decrypted_path(up, wd)
        dst = os.path.join(wd, f"{secrets.token_hex(4)}.docx")
        doc = open_pdf(up)
        scanned = sum(len(p.get_text().strip()) for p in doc) < 20 * doc.page_count
        doc.close()
        if scanned and _b(opts.get("ocr"), True) and tesseract_bin():
            ocr_res = ocr([Upload(up.name, src)], {"languages": opts.get("languages") or ["eng"],
                                                  "skip_text": True}, wd)
            src = os.path.join(wd, f"ocr_{secrets.token_hex(4)}.pdf")
            Path(src).write_bytes(ocr_res.files[0][1])
        try:
            cv = Converter(src)
            cv.convert(dst, multi_processing=False)
            cv.close()
        except Exception as e:
            if Path(dst).exists() and os.path.getsize(dst) > 0:
                pass
            else:
                raise UserError(f"Could not convert “{up.name}” to Word: {e}")
        out.append((f"{up.stem}.docx", Path(dst).read_bytes()))
    return Result(out)


def pdf_to_powerpoint(files, opts, wd):
    from pptx import Presentation
    from pptx.util import Emu
    pdfs = require_pdfs(files, 1)
    out = []
    for up in pdfs:
        doc = open_pdf(up)
        prs = Presentation()
        r0 = doc[0].rect
        w_emu = 12192000  # 13.333 in
        h_emu = int(w_emu * r0.height / r0.width)
        if h_emu > 51206400:
            h_emu = 51206400
            w_emu = int(h_emu * r0.width / r0.height)
        prs.slide_width, prs.slide_height = Emu(w_emu), Emu(max(914400, h_emu))
        blank = prs.slide_layouts[6]
        for page in doc:
            slide = prs.slides.add_slide(blank)
            pix = page.get_pixmap(dpi=170, alpha=False)
            img = io.BytesIO(pix.tobytes("jpg", jpg_quality=90))
            sw, sh = prs.slide_width, prs.slide_height
            pr = page.rect.width / page.rect.height
            if sw / sh > pr:
                h = sh
                w = int(h * pr)
            else:
                w = sw
                h = int(w / pr)
            slide.shapes.add_picture(img, int((sw - w) / 2), int((sh - h) / 2), w, h)
            text = page.get_text().strip()
            if text:
                slide.notes_slide.notes_text_frame.text = text[:20000]
        buf = io.BytesIO()
        prs.save(buf)
        out.append((f"{up.stem}.pptx", buf.getvalue()))
        doc.close()
    return Result(out)


def _cell(v):
    if v is None:
        return None
    s = str(v).strip()
    if re.fullmatch(r"-?\d{1,15}", s.replace(",", "")) and not (len(s) > 1 and s.startswith("0")):
        try:
            return int(s.replace(",", ""))
        except ValueError:
            return s
    if re.fullmatch(r"-?\d[\d,]*\.\d+", s):
        try:
            return float(s.replace(",", ""))
        except ValueError:
            return s
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", s)


def _text_rows(page):
    words = page.get_text("words", sort=True)
    lines = {}
    for w in words:
        key = (w[5], w[6])
        lines.setdefault(key, []).append(w)
    rows = []
    for key in sorted(lines, key=lambda k: (min(x[1] for x in lines[k]), k)):
        ws = sorted(lines[key], key=lambda x: x[0])
        cells, cur, last = [], [], None
        for w in ws:
            avg = (w[2] - w[0]) / max(1, len(w[4]))
            if last is not None and w[0] - last > max(8, avg * 2.2):
                cells.append(" ".join(cur))
                cur = []
            cur.append(w[4])
            last = w[2]
        if cur:
            cells.append(" ".join(cur))
        rows.append(cells)
    return rows


def pdf_to_excel(files, opts, wd):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.styles import Font
    pdfs = require_pdfs(files, 1)
    out = []
    for up in pdfs:
        doc = open_pdf(up)
        wb = Workbook()
        wb.remove(wb.active)
        one_sheet = opts.get("layout", "page") == "single"
        single = wb.create_sheet("Data") if one_sheet else None
        tables_found = 0
        for page in doc:
            rows_blocks = []
            try:
                tabs = page.find_tables()
                for t in tabs.tables:
                    data = t.extract()
                    if data and any(any(c for c in r) for r in data):
                        rows_blocks.append(data)
            except Exception:
                pass
            tables_found += len(rows_blocks)
            if not rows_blocks:
                tr = _text_rows(page)
                if tr:
                    rows_blocks = [tr]
            if not rows_blocks:
                continue
            ws = single or wb.create_sheet(f"Page {page.number + 1}")
            for bi, block in enumerate(rows_blocks):
                if ws.max_row > 1 or bi:
                    ws.append([])
                for ri, row in enumerate(block):
                    ws.append([_cell(c) for c in row])
                    if ri == 0 and len(block) > 1 and tables_found:
                        for c in ws[ws.max_row]:
                            c.font = Font(bold=True)
            for col in range(1, ws.max_column + 1):
                width = max((len(str(c.value)) for c in ws[get_column_letter(col)] if c.value is not None), default=8)
                ws.column_dimensions[get_column_letter(col)].width = min(60, max(8, width + 2))
        if not wb.sheetnames:
            raise UserError(f"No text or tables found in “{up.name}”. If it is a scanned PDF, run OCR PDF first.")
        buf = io.BytesIO()
        wb.save(buf)
        out.append((f"{up.stem}.xlsx", buf.getvalue()))
        doc.close()
    return Result(out)


def _srgb_icc():
    from PIL import ImageCms
    return ImageCms.ImageCmsProfile(ImageCms.createProfile("sRGB")).tobytes()


def pdf_to_pdfa(files, opts, wd):
    gs = gs_bin()
    if not gs:
        raise UserError("PDF/A conversion needs Ghostscript (free). Install it (see README) and restart the app.")
    pdfs = require_pdfs(files, 1)
    level = str(opts.get("level", "2b")).lower()
    part = {"1b": 1, "2b": 2, "3b": 3}.get(level, 2)
    icc = os.path.join(wd, "srgb.icc")
    Path(icc).write_bytes(_srgb_icc())
    icc_ps = icc.replace("\\", "/").replace("(", "\\(").replace(")", "\\)")
    defs = os.path.join(wd, "pdfa_def.ps")
    Path(defs).write_text(f"""%!
/ICCProfile ({icc_ps}) def
[/_objdef {{icc_PDFA}} /type /stream /OBJ pdfmark
[{{icc_PDFA}} <</N 3>> /PUT pdfmark
[{{icc_PDFA}} ICCProfile (r) file /PUT pdfmark
[/_objdef {{OutputIntent_PDFA}} /type /dict /OBJ pdfmark
[{{OutputIntent_PDFA}} <</Type /OutputIntent /S /GTS_PDFA1 /DestOutputProfile {{icc_PDFA}} /OutputConditionIdentifier (sRGB) /Info (sRGB IEC61966-2.1)>> /PUT pdfmark
[{{Catalog}} <</OutputIntents [ {{OutputIntent_PDFA}} ]>> /PUT pdfmark
""")
    out = []
    for up in pdfs:
        src = decrypted_path(up, wd)
        dst = os.path.join(wd, f"pdfa_{secrets.token_hex(4)}.pdf")
        cmd = [gs, f"-dPDFA={part}", "-dBATCH", "-dNOPAUSE", "-dNOSAFER", "-dNOOUTERSAVE",
               "-dPDFACompatibilityPolicy=1", "-sColorConversionStrategy=RGB", "-sProcessColorModel=DeviceRGB",
               "-sDEVICE=pdfwrite", f"-dCompatibilityLevel={'1.4' if part == 1 else '1.7'}",
               f"-sOutputFile={dst}", defs, src]
        proc = run(cmd, timeout=600)
        if not os.path.exists(dst) or os.path.getsize(dst) == 0:
            raise UserError("PDF/A conversion failed: " + (proc.stderr or b"").decode("utf-8", "ignore")[:300])
        out.append((f"{up.stem}_pdfa.pdf", Path(dst).read_bytes()))
    return Result(out, {"level": f"PDF/A-{part}b"})


def pdf_to_markdown(files, opts, wd):
    pdfs = require_pdfs(files, 1)
    out, preview = [], ""
    for up in pdfs:
        doc = open_pdf(up)
        md = None
        try:
            import pymupdf4llm
            md = pymupdf4llm.to_markdown(doc, show_progress=False)
        except Exception:
            md = None
        if not md or not md.strip():
            md = _simple_markdown(doc)
        if not md.strip():
            raise UserError(f"No text found in “{up.name}”. If it is scanned, run OCR PDF first.")
        preview = preview or md[:20000]
        out.append((f"{up.stem}.md", md.encode("utf-8")))
        doc.close()
    return Result(out, {"preview": preview})


def _simple_markdown(doc):
    sizes = Counter()
    for page in doc:
        for b in page.get_text("dict")["blocks"]:
            for l in b.get("lines", []):
                for s in l["spans"]:
                    sizes[round(s["size"])] += len(s["text"])
    body = sizes.most_common(1)[0][0] if sizes else 11
    md = []
    for page in doc:
        for b in page.get_text("dict")["blocks"]:
            if b["type"] != 0:
                continue
            text = " ".join("".join(s["text"] for s in l["spans"]) for l in b["lines"]).strip()
            if not text:
                continue
            size = max((s["size"] for l in b["lines"] for s in l["spans"]), default=body)
            if size >= body * 1.6:
                md.append("# " + text)
            elif size >= body * 1.3:
                md.append("## " + text)
            elif size >= body * 1.12:
                md.append("### " + text)
            elif re.match(r"^[•●▪\-–*]\s*", text):
                md.append("- " + re.sub(r"^[•●▪\-–*]\s*", "", text))
            else:
                md.append(text)
        md.append("\n---\n")
    return "\n\n".join(md)


# =========================================================================== #
# EDIT
# =========================================================================== #
def _pages_opt(opts, doc):
    return parse_pages(opts.get("pages"), doc.page_count)


def page_numbers(files, opts, wd):
    pdfs = require_pdfs(files, 1)
    pos = opts.get("position", "bc")
    margin = {"small": 18, "recommended": 32, "big": 54}.get(opts.get("margin", "recommended"), 32)
    first = int(_num(opts.get("first_number"), 1, -100000, 10 ** 7, int))
    fmt = opts.get("format") or "{n}"
    if fmt == "custom":
        fmt = opts.get("custom") or "{n}"
    size = _num(opts.get("size"), 11, 4, 200)
    color = rgb_to_hex(hex_to_rgb(opts.get("color"), (0, 0, 0)))
    family = opts.get("font", "sans")
    bold = _b(opts.get("bold"))
    facing = opts.get("mode") == "facing"
    out = []
    for up in pdfs:
        doc = open_pdf(up)
        pages = _pages_opt(opts, doc)
        start_from = int(_num(opts.get("start_page"), 1, 1, doc.page_count, int)) - 1
        pages = [p for p in pages if p >= start_from]
        total = first + len(pages) - 1
        for i, p in enumerate(pages):
            page = doc[p]
            normalize_page(page)
            r = page.rect
            vert, horiz = pos[0], pos[1]
            if facing and (p % 2 == 1) and horiz in "lr":
                horiz = "r" if horiz == "l" else "l"
            align = {"l": "left", "c": "center", "r": "right"}[horiz]
            h = size * 1.6
            y0 = margin if vert == "t" else r.height - margin - h
            if vert == "m":
                y0 = (r.height - h) / 2
            box = pymupdf.Rect(margin, y0, r.width - margin, y0 + h)
            label = fmt.replace("{n}", str(first + i)).replace("{total}", str(total)).replace("{p}", str(p + 1))
            insert_text_box(page, box, label, size=size, color=color, family=family, bold=bold, align=align)
        out.append((f"{up.stem}_numbered.pdf", pdf_bytes(doc)))
        doc.close()
    return Result(out)


def _grid_rect(page_rect, w, h, position, margin=24):
    W, H = page_rect.width, page_rect.height
    v, hz = position[0], position[1]
    x = {"l": margin, "c": (W - w) / 2, "r": W - w - margin}[hz]
    y = {"t": margin, "m": (H - h) / 2, "b": H - h - margin}[v]
    return pymupdf.Rect(x, y, x + w, y + h)


def _rotated_size(w, h, deg):
    a = math.radians(deg)
    return abs(w * math.cos(a)) + abs(h * math.sin(a)), abs(w * math.sin(a)) + abs(h * math.cos(a))


def watermark(files, opts, wd, image_bytes=None):
    pdfs = [f for f in files if f.ext == ".pdf"]
    imgs = [f for f in files if f.ext in IMAGE_EXTS]
    if not pdfs:
        raise UserError("Please select at least one PDF.")
    wtype = opts.get("type", "text")
    position = opts.get("position", "mc")
    mosaic = position == "mosaic"
    angle = _num(opts.get("rotation"), 0, -360, 360)
    opacity = _num(opts.get("opacity"), 0.5, 0.0, 1.0)
    below = opts.get("layer") == "below"
    if wtype == "image":
        if not imgs and not image_bytes:
            raise UserError("Please choose a watermark image.")
        src = Image.open(io.BytesIO(image_bytes) if image_bytes else imgs[0].path)
        src = ImageOps.exif_transpose(src).convert("RGBA")
        if opacity < 1:
            alpha = src.split()[-1].point(lambda a: int(a * opacity))
            src.putalpha(alpha)
        if angle % 360:
            src = src.rotate(angle, expand=True, resample=Image.BICUBIC)
        buf = io.BytesIO()
        src.save(buf, "PNG")
        stamp_png = buf.getvalue()
        sw, sh = src.size
        scale_pct = _num(opts.get("scale"), 40, 1, 100) / 100
    else:
        text = (opts.get("text") or "").strip()
        if not text:
            raise UserError("Please enter the watermark text.")
        size = _num(opts.get("size"), 48, 4, 400)
        stamp_doc, clip = make_text_stamp(text, size=size, color=rgb_to_hex(hex_to_rgb(opts.get("color"), (1, 0, 0))),
                                          family=opts.get("font", "sans"), bold=_b(opts.get("bold"), True),
                                          italic=_b(opts.get("italic")), opacity=opacity)
        sw, sh = clip.width, clip.height
    out = []
    for up in pdfs:
        doc = open_pdf(up)
        for p in _pages_opt(opts, doc):
            page = doc[p]
            normalize_page(page)
            pr = page.rect
            if wtype == "image":
                w = pr.width * scale_pct
                h = w * sh / sw
                if h > pr.height * 0.95:
                    h = pr.height * 0.95
                    w = h * sw / sh
                bw, bh = w, h
            else:
                bw, bh = _rotated_size(sw, sh, angle)
                f = min(1.0, pr.width * 0.95 / bw, pr.height * 0.95 / bh)
                bw, bh = bw * f, bh * f
            rects = []
            if mosaic:
                gx, gy = bw * 1.4, bh * 1.8
                cols = max(1, int(pr.width // gx) + 1)
                rows_ = max(1, int(pr.height // gy) + 1)
                ox = (pr.width - (cols - 1) * gx - bw) / 2
                oy = (pr.height - (rows_ - 1) * gy - bh) / 2
                for i in range(cols):
                    for j in range(rows_):
                        x, y = ox + i * gx, oy + j * gy
                        rects.append(pymupdf.Rect(x, y, x + bw, y + bh))
            else:
                rects.append(_grid_rect(pr, bw, bh, position if len(position) == 2 else "mc"))
            for rect in rects:
                if wtype == "image":
                    page.insert_image(rect, stream=stamp_png, keep_proportion=True, overlay=not below)
                else:
                    page.show_pdf_page(rect, stamp_doc, 0, clip=clip, rotate=angle, overlay=not below,
                                       keep_proportion=True)
        out.append((f"{up.stem}_watermarked.pdf", pdf_bytes(doc)))
        doc.close()
    return Result(out)


def crop(files, opts, wd):
    pdfs = require_pdfs(files, 1)
    mode = opts.get("mode", "margins")
    out = []
    for up in pdfs:
        doc = open_pdf(up)
        if mode == "box":
            box = opts.get("box") or {}
            fx0, fy0, fx1, fy1 = (_num(box.get(k), d, 0, 1) for k, d in
                                  (("x0", 0), ("y0", 0), ("x1", 1), ("y1", 1)))
            if fx1 - fx0 < 0.01 or fy1 - fy0 < 0.01:
                raise UserError("The crop area is too small.")
            if opts.get("apply") == "current":
                pages = [int(_num(opts.get("page"), 0, 0, doc.page_count - 1, int))]
            else:
                pages = _pages_opt(opts, doc)
        else:
            m = {k: _num(opts.get(k), 0, 0, 5000) for k in ("top", "right", "bottom", "left")}
            pages = _pages_opt(opts, doc)
        for p in pages:
            page = doc[p]
            normalize_page(page)
            cb = page.cropbox
            if mode == "box":
                new = pymupdf.Rect(cb.x0 + fx0 * cb.width, cb.y0 + fy0 * cb.height,
                                   cb.x0 + fx1 * cb.width, cb.y0 + fy1 * cb.height)
            else:
                new = pymupdf.Rect(cb.x0 + m["left"], cb.y0 + m["top"], cb.x1 - m["right"], cb.y1 - m["bottom"])
            new = new & page.mediabox
            if new.width < 10 or new.height < 10:
                raise UserError(f"Margins are too large for page {p + 1}.")
            page.set_cropbox(new)
        if _b(opts.get("permanent"), False):
            # make crop permanent by rewriting mediabox too
            for p in pages:
                page = doc[p]
                page.set_mediabox(page.cropbox)
        out.append((f"{up.stem}_cropped.pdf", pdf_bytes(doc)))
        doc.close()
    return Result(out)


def _data_url_bytes(s):
    if not s:
        return None
    if "," in s and s.startswith("data:"):
        s = s.split(",", 1)[1]
    try:
        return base64.b64decode(s)
    except Exception:
        raise UserError("An embedded image is invalid.")


def edit(files, opts, wd):
    """opts.items: [{page, type, x, y, w, h, ...}] – coordinates as fractions (0..1) of the visible page."""
    up = require_pdfs(files, 1, 1)[0]
    doc = open_pdf(up)
    items = opts.get("items") or []
    if not items:
        raise UserError("Nothing to apply – add some text, images, shapes or a signature first.")
    touched = set()
    for it in items:
        p = int(it.get("page", 0))
        if p < 0 or p >= doc.page_count:
            continue
        page = doc[p]
        if p not in touched:
            normalize_page(page)
            touched.add(p)
        W, H = page.rect.width, page.rect.height
        x, y = _num(it.get("x"), 0) * W, _num(it.get("y"), 0) * H
        w, h = _num(it.get("w"), 0.1) * W, _num(it.get("h"), 0.05) * H
        rect = pymupdf.Rect(x, y, x + max(w, 1), y + max(h, 1))
        t = it.get("type")
        opacity = _num(it.get("opacity"), 1, 0, 1)
        if t == "text":
            text = str(it.get("text", ""))
            if not text.strip():
                continue
            if it.get("bg"):
                page.draw_rect(rect, color=None, fill=hex_to_rgb(it["bg"], (1, 1, 1)), overlay=True)
            insert_text_box(page, rect, text, size=_num(it.get("size"), 14, 2, 400),
                            color=rgb_to_hex(hex_to_rgb(it.get("color"))), family=it.get("font", "sans"),
                            bold=_b(it.get("bold")), italic=_b(it.get("italic")), align=it.get("align", "left"))
        elif t in ("image", "signature"):
            data = _data_url_bytes(it.get("src"))
            if not data:
                continue
            if opacity < 1:
                im = Image.open(io.BytesIO(data)).convert("RGBA")
                im.putalpha(im.split()[-1].point(lambda a: int(a * opacity)))
                b = io.BytesIO()
                im.save(b, "PNG")
                data = b.getvalue()
            page.insert_image(rect, stream=data, keep_proportion=False, overlay=True)
        elif t in ("rect", "ellipse", "whiteout", "highlight"):
            stroke = None if not it.get("stroke") or t in ("whiteout", "highlight") else hex_to_rgb(it.get("stroke"))
            fill = hex_to_rgb(it.get("fill")) if it.get("fill") else None
            if t == "whiteout":
                fill, opacity = (1, 1, 1), 1
            if t == "highlight":
                fill = hex_to_rgb(it.get("fill") or "#ffeb3b")
                opacity = _num(it.get("opacity"), 0.4, 0, 1)
            shape = page.new_shape()
            if t == "ellipse":
                shape.draw_oval(rect)
            else:
                shape.draw_rect(rect)
            shape.finish(color=stroke, fill=fill, width=_num(it.get("width"), 2, 0, 50),
                         fill_opacity=opacity, stroke_opacity=opacity if stroke else 1)
            shape.commit(overlay=True)
        elif t in ("line", "arrow"):
            p1 = pymupdf.Point(_num(it.get("x1"), 0) * W, _num(it.get("y1"), 0) * H)
            p2 = pymupdf.Point(_num(it.get("x2"), 0) * W, _num(it.get("y2"), 0) * H)
            col = hex_to_rgb(it.get("stroke") or it.get("color"))
            lw = _num(it.get("width"), 2, 0.2, 50)
            shape = page.new_shape()
            shape.draw_line(p1, p2)
            if t == "arrow" and abs(p2 - p1) > 1:
                ang = math.atan2(p2.y - p1.y, p2.x - p1.x)
                L = max(8, lw * 4)
                for d in (2.6, -2.6):
                    shape.draw_line(p2, pymupdf.Point(p2.x + L * math.cos(ang + d), p2.y + L * math.sin(ang + d)))
            shape.finish(color=col, width=lw, stroke_opacity=opacity, lineCap=1)
            shape.commit(overlay=True)
        elif t == "draw":
            pts = [pymupdf.Point(_num(a, 0) * W, _num(b, 0) * H) for a, b in (it.get("points") or [])]
            if len(pts) < 2:
                continue
            shape = page.new_shape()
            shape.draw_polyline(pts)
            shape.finish(color=hex_to_rgb(it.get("stroke") or it.get("color")), width=_num(it.get("width"), 2, 0.2, 50),
                         closePath=False, stroke_opacity=opacity, lineCap=1, lineJoin=1)
            shape.commit(overlay=True)
    suffix = "signed" if opts.get("mode") == "sign" else "edited"
    return Result([(f"{up.stem}_{suffix}.pdf", pdf_bytes(doc))])


# =========================================================================== #
# FORMS
# =========================================================================== #
WTYPES = {
    pymupdf.PDF_WIDGET_TYPE_TEXT: "text", pymupdf.PDF_WIDGET_TYPE_CHECKBOX: "checkbox",
    pymupdf.PDF_WIDGET_TYPE_COMBOBOX: "dropdown", pymupdf.PDF_WIDGET_TYPE_LISTBOX: "listbox",
    pymupdf.PDF_WIDGET_TYPE_RADIOBUTTON: "radio", pymupdf.PDF_WIDGET_TYPE_SIGNATURE: "signature",
    pymupdf.PDF_WIDGET_TYPE_BUTTON: "button",
}


def form_fields(files, opts, wd):
    up = require_pdfs(files, 1, 1)[0]
    doc = open_pdf(up)
    fields = []
    for page in doc:
        r = page.rect
        for w in page.widgets() or []:
            typ = WTYPES.get(w.field_type, "other")
            wr = w.rect * page.rotation_matrix if page.rotation else w.rect
            val = w.field_value
            if typ in ("checkbox", "radio"):
                on = w.on_state()
                val = bool(val) and str(val) not in ("Off", "False", "false", "") and (on is True or str(val) == str(on) or val is True)
            fields.append({
                "id": str(w.xref), "page": page.number, "name": w.field_name, "label": w.field_label or "",
                "type": typ, "value": val if isinstance(val, bool) else ("" if val is None else str(val)),
                "options": list(w.choice_values or []), "multiline": bool(w.field_flags & 4096) if typ == "text" else False,
                "x": wr.x0 / r.width, "y": wr.y0 / r.height, "w": wr.width / r.width, "h": wr.height / r.height,
                "readonly": bool(w.field_flags & 1),
            })
    return Result(data={"fields": fields, "detected": _detect_fields(doc) if not fields else []})


def _detect_fields(doc):
    found = []
    for page in doc:
        r = page.rect
        if page.rotation:
            continue
        for wd_ in page.get_text("words"):
            txt = wd_[4]
            m = re.search(r"_{4,}", txt)
            if m:
                x0, y0, x1, y1 = wd_[:4]
                cw = (x1 - x0) / max(1, len(txt))
                fx0 = x0 + m.start() * cw
                fx1 = x0 + m.end() * cw
                h = max(12, (y1 - y0) * 1.1)
                found.append({"page": page.number, "type": "text", "x": fx0 / r.width, "y": (y1 - h) / r.height,
                              "w": (fx1 - fx0) / r.width, "h": h / r.height})
            elif txt in ("☐", "□", "❏", "▢", "◻"):
                x0, y0, x1, y1 = wd_[:4]
                found.append({"page": page.number, "type": "checkbox", "x": x0 / r.width, "y": y0 / r.height,
                              "w": (x1 - x0) / r.width, "h": (y1 - y0) / r.height})
        try:
            for d in page.get_drawings():
                rr = d.get("rect")
                if not rr or d.get("fill") not in (None, (1.0, 1.0, 1.0)):
                    continue
                items = d.get("items", [])
                if len(items) != 1 or items[0][0] not in ("re", "qu"):
                    continue
                if 7 <= rr.width <= 22 and 7 <= rr.height <= 22 and abs(rr.width - rr.height) < 3:
                    found.append({"page": page.number, "type": "checkbox", "x": rr.x0 / r.width, "y": rr.y0 / r.height,
                                  "w": rr.width / r.width, "h": rr.height / r.height})
                elif rr.width > 60 and 12 <= rr.height <= 40 and rr.width < r.width * 0.9:
                    if not page.get_text("text", clip=rr).strip():
                        found.append({"page": page.number, "type": "text", "x": rr.x0 / r.width, "y": rr.y0 / r.height,
                                      "w": rr.width / r.width, "h": rr.height / r.height})
        except Exception:
            pass
    return found[:500]


def forms(files, opts, wd):
    """opts.values: {xref: value} to fill existing fields; opts.new_fields: [...] to create; opts.flatten"""
    up = require_pdfs(files, 1, 1)[0]
    doc = open_pdf(up)
    values = opts.get("values") or {}
    new_fields = opts.get("new_fields") or []
    filled = created = 0
    for page in doc:
        for w in page.widgets() or []:
            key = str(w.xref)
            if key not in values:
                continue
            v = values[key]
            try:
                if w.field_type in (pymupdf.PDF_WIDGET_TYPE_CHECKBOX, pymupdf.PDF_WIDGET_TYPE_RADIOBUTTON):
                    w.field_value = w.on_state() if _b(v) else "Off"
                elif w.field_type in (pymupdf.PDF_WIDGET_TYPE_BUTTON, pymupdf.PDF_WIDGET_TYPE_SIGNATURE):
                    continue
                else:
                    w.field_value = "" if v is None else str(v)
                w.update()
                filled += 1
            except Exception:
                continue
    names = {w.field_name for p in doc for w in (p.widgets() or [])}
    for i, f in enumerate(new_fields):
        p = int(f.get("page", 0))
        if p < 0 or p >= doc.page_count:
            continue
        page = doc[p]
        normalize_page(page)
        W, H = page.rect.width, page.rect.height
        rect = pymupdf.Rect(_num(f.get("x"), 0) * W, _num(f.get("y"), 0) * H,
                            (_num(f.get("x"), 0) + _num(f.get("w"), 0.2)) * W,
                            (_num(f.get("y"), 0) + _num(f.get("h"), 0.03)) * H)
        base = re.sub(r"[^\w\-]", "_", str(f.get("name") or f"{f.get('type', 'field')}_{i + 1}"))[:60] or f"field_{i + 1}"
        name, k = base, 2
        while name in names:
            name, k = f"{base}_{k}", k + 1
        names.add(name)
        w = pymupdf.Widget()
        w.field_name = name
        w.rect = rect
        w.border_color = (0.35, 0.35, 0.35)
        w.border_width = 0.8
        w.fill_color = (0.93, 0.95, 1.0)
        typ = f.get("type", "text")
        if typ == "checkbox":
            w.field_type = pymupdf.PDF_WIDGET_TYPE_CHECKBOX
            w.field_value = _b(f.get("value"))
        elif typ in ("dropdown", "listbox"):
            w.field_type = pymupdf.PDF_WIDGET_TYPE_COMBOBOX if typ == "dropdown" else pymupdf.PDF_WIDGET_TYPE_LISTBOX
            opts_list = [o.strip() for o in (f.get("options") or []) if str(o).strip()] or ["Option 1", "Option 2"]
            w.choice_values = opts_list
            w.field_value = f.get("value") if f.get("value") in opts_list else opts_list[0]
            w.text_fontsize = 0
        else:
            w.field_type = pymupdf.PDF_WIDGET_TYPE_TEXT
            w.field_value = str(f.get("value") or "")
            w.text_fontsize = 0
            if typ == "multiline":
                w.field_flags |= pymupdf.PDF_TX_FIELD_IS_MULTILINE
        if _b(f.get("required")):
            w.field_flags |= pymupdf.PDF_FIELD_IS_REQUIRED
        page.add_widget(w)
        created += 1
    if not filled and not created:
        raise UserError("No form changes to save – fill in a field or add new fields first.")
    if _b(opts.get("flatten")):
        doc.bake(annots=False, widgets=True)
    return Result([(f"{up.stem}_form.pdf", pdf_bytes(doc))], {"filled": filled, "created": created})


# =========================================================================== #
# SECURITY
# =========================================================================== #
def unlock(files, opts, wd):
    pdfs = require_pdfs(files, 1)
    out = []
    for up in pdfs:
        doc = pymupdf.open(up.path, filetype="pdf") if _safe_open(up) else None
        if doc.needs_pass:
            pw = up.password or opts.get("password")
            if not pw:
                raise UserError(f"“{up.name}” needs a password to be unlocked. Please enter it.")
            if not doc.authenticate(pw):
                raise UserError(f"Wrong password for “{up.name}”.")
        out.append((f"{up.stem}_unlocked.pdf", doc.tobytes(garbage=3, deflate=True, encryption=pymupdf.PDF_ENCRYPT_NONE)))
        doc.close()
    return Result(out)


def _safe_open(up):
    try:
        pymupdf.open(up.path, filetype="pdf").close()
        return True
    except Exception:
        raise UserError(f"“{up.name}” is not a valid PDF.")


def protect(files, opts, wd):
    pdfs = require_pdfs(files, 1)
    pw = opts.get("password") or ""
    if len(pw) < 1:
        raise UserError("Please enter a password.")
    if opts.get("confirm") is not None and opts.get("confirm") != pw:
        raise UserError("Passwords do not match.")
    owner = opts.get("owner_password") or secrets.token_urlsafe(16)
    perm = 0
    if _b(opts.get("allow_print"), True):
        perm |= pymupdf.PDF_PERM_PRINT | pymupdf.PDF_PERM_PRINT_HQ
    if _b(opts.get("allow_copy"), True):
        perm |= pymupdf.PDF_PERM_COPY | pymupdf.PDF_PERM_ACCESSIBILITY
    if _b(opts.get("allow_modify"), True):
        perm |= pymupdf.PDF_PERM_MODIFY | pymupdf.PDF_PERM_ASSEMBLE
    if _b(opts.get("allow_annotate"), True):
        perm |= pymupdf.PDF_PERM_ANNOTATE | pymupdf.PDF_PERM_FORM
    out = []
    for up in pdfs:
        doc = open_pdf(up)
        data = doc.tobytes(garbage=3, deflate=True, encryption=pymupdf.PDF_ENCRYPT_AES_256,
                           user_pw=pw, owner_pw=owner, permissions=perm)
        out.append((f"{up.stem}_protected.pdf", data))
        doc.close()
    return Result(out)


REDACT_PATTERNS = {
    "email": r"[\w.+-]+@[\w-]+\.[\w.-]+",
    "phone": r"(?<!\w)\+?\d[\d\s().-]{7,}\d(?!\w)",
    "url": r"https?://\S+|www\.\S+",
    "date": r"\b\d{1,4}[/.-]\d{1,2}[/.-]\d{1,4}\b",
    "card": r"\b(?:\d[ -]?){13,19}\b",
    "number": r"\b\d{4,}\b",
}


def redact(files, opts, wd):
    pdfs = require_pdfs(files, 1)
    terms = opts.get("terms") or []
    if isinstance(terms, str):
        terms = [t for t in re.split(r"[\n,]+", terms)]
    terms = [t.strip() for t in terms if t and t.strip()]
    patterns = [REDACT_PATTERNS[p] for p in (opts.get("patterns") or []) if p in REDACT_PATTERNS]
    areas = opts.get("areas") or []
    fill = hex_to_rgb(opts.get("color") or "#000000")
    if not terms and not patterns and not areas:
        raise UserError("Add words to search, choose a pattern or mark areas to redact.")
    out, total = [], 0
    for fi, up in enumerate(pdfs):
        doc = open_pdf(up)
        count = 0
        for page in doc:
            normalize_page(page)
            hits = []
            text = page.get_text() if patterns else ""
            words = set(terms)
            for pat in patterns:
                for m in re.finditer(pat, text):
                    s = m.group(0).strip()
                    if len(s) >= 3:
                        words.add(s)
            for t in words:
                hits.extend(page.search_for(t, quads=False))
            W, H = page.rect.width, page.rect.height
            for a in areas:
                if int(a.get("page", -1)) == page.number and int(a.get("file", 0)) == fi:
                    hits.append(pymupdf.Rect(_num(a.get("x0"), 0) * W, _num(a.get("y0"), 0) * H,
                                             _num(a.get("x1"), 0) * W, _num(a.get("y1"), 0) * H))
            for r in hits:
                if r.is_empty or r.is_infinite:
                    continue
                page.add_redact_annot(r, fill=fill, cross_out=False)
                count += 1
            if hits:
                page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_PIXELS,
                                      graphics=pymupdf.PDF_REDACT_LINE_ART_REMOVE_IF_TOUCHED)
        if _b(opts.get("metadata"), True):
            doc.set_metadata({})
            doc.del_xml_metadata()
        total += count
        out.append((f"{up.stem}_redacted.pdf", pdf_bytes(doc, garbage=4)))
        doc.close()
    if total == 0 and not areas:
        raise UserError("None of the search words or patterns were found – nothing was redacted.")
    return Result(out, {"redactions": total})


# =========================================================================== #
# COMPARE
# =========================================================================== #
def compare(files, opts, wd):
    pdfs = require_pdfs(files, 2, 2)
    a, b = open_pdf(pdfs[0]), open_pdf(pdfs[1])

    def words(doc):
        res = []
        for page in doc:
            for w in page.get_text("words", sort=True):
                res.append((w[4], page.number, pymupdf.Rect(w[:4])))
        return res

    wa, wb = words(a), words(b)
    if len(wa) + len(wb) > 400000:
        raise UserError("Documents are too large to compare (max ~200,000 words each).")
    sm = difflib.SequenceMatcher(None, [w[0] for w in wa], [w[0] for w in wb], autojunk=False)
    ops, dels, ins = [], [], []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        ta = " ".join(w[0] for w in wa[i1:i2])
        tb = " ".join(w[0] for w in wb[j1:j2])
        ops.append({"op": tag, "a": ta[:5000], "b": tb[:5000]})
        if tag in ("delete", "replace"):
            dels.extend(wa[i1:i2])
        if tag in ("insert", "replace"):
            ins.extend(wb[j1:j2])
    for doc, lst, color in ((a, dels, (1, 0.45, 0.45)), (b, ins, (0.4, 0.85, 0.4))):
        by_page = {}
        for text, pno, r in lst:
            by_page.setdefault(pno, []).append(r)
        for pno, rects in by_page.items():
            page = doc[pno]
            for r in rects:
                annot = page.add_highlight_annot(r)
                annot.set_colors(stroke=color)
                annot.update()
    # side-by-side report
    report = pymupdf.open()
    n = max(a.page_count, b.page_count)
    for i in range(n):
        ra = a[i].rect if i < a.page_count else a[-1].rect
        rb = b[i].rect if i < b.page_count else b[-1].rect
        h = max(ra.height, rb.height) + 40
        page = report.new_page(width=ra.width + rb.width + 30, height=h)
        insert_text_box(page, pymupdf.Rect(10, 8, ra.width, 30), f"{pdfs[0].name} – page {i + 1}", size=10, color="#b71c1c")
        insert_text_box(page, pymupdf.Rect(ra.width + 20, 8, ra.width + 20 + rb.width, 30),
                        f"{pdfs[1].name} – page {i + 1}", size=10, color="#1b5e20")
        if i < a.page_count:
            page.show_pdf_page(pymupdf.Rect(10, 32, 10 + ra.width, 32 + ra.height), _flatten_annots(a, i), 0)
        if i < b.page_count:
            page.show_pdf_page(pymupdf.Rect(ra.width + 20, 32, ra.width + 20 + rb.width, 32 + rb.height),
                               _flatten_annots(b, i), 0)
    changes = sum(1 for o in ops if o["op"] != "equal")
    ratio = round(sm.ratio() * 100, 1) if (wa or wb) else 100.0
    data = {"ops": ops if len(ops) < 20000 else ops[:20000], "changes": changes, "similarity": ratio,
            "deleted_words": len(dels), "inserted_words": len(ins),
            "report": base64.b64encode(pdf_bytes(report)).decode(),
            "report_name": f"compare_{pdfs[0].stem}_vs_{pdfs[1].stem}.pdf"}
    return Result(data=data)


def _flatten_annots(doc, pno):
    tmp = pymupdf.open()
    tmp.insert_pdf(doc, from_page=pno, to_page=pno, annots=True)
    tmp.bake(annots=True, widgets=True)
    return pymupdf.open("pdf", tmp.tobytes())


# =========================================================================== #
# AI-like tools: summarize, translate
# =========================================================================== #
STOP = set("""a an the and or but if then else of to in on at by for with about against between into through during
before after above below from up down out off over under again further once here there when where why how all any both
each few more most other some such no nor not only own same so than too very s t can will just don should now is are was
were be been being have has had having do does did doing i me my we our you your he him his she her it its they them
their what which who whom this that these those am would could also may might must shall et al fig table""".split())


def _sentences(text):
    text = re.sub(r"-\n(\w)", r"\1", text)
    text = re.sub(r"[ \t]*\n[ \t]*", " ", text)
    text = re.sub(r"\s{2,}", " ", text)
    parts = re.split(r"(?<=[.!?।॥。！？])\s+(?=[\"'“(\[]?[A-Z0-9\u0980-\u09FF\u0900-\u097F\u0600-\u06FF\u4e00-\u9fff])", text)
    seen, out = set(), []
    for p in parts:
        p = p.strip()
        key = re.sub(r"\W+", "", p.lower())
        if len(p) > 20 and key not in seen:
            seen.add(key)
            out.append(p)
    return out


def summarize_text(text, length="medium"):
    sents = _sentences(text)
    if not sents:
        raise UserError("Not enough text to summarize. If this is a scanned PDF, run OCR PDF first.")
    words_all = [w.lower() for w in re.findall(r"\w+", text) if len(w) > 2]
    freq = Counter(w for w in words_all if w not in STOP and not w.isdigit())
    if not freq:
        raise UserError("Not enough text to summarize.")
    top = freq.most_common(1)[0][1]
    ratio = {"short": 0.1, "medium": 0.2, "long": 0.35}.get(length, 0.2)
    k = max(3, min(40 if length == "long" else 25, int(len(sents) * ratio) or 3))
    scored = []
    for i, s in enumerate(sents):
        ws = [w.lower() for w in re.findall(r"\w+", s)]
        if not ws:
            continue
        score = sum(freq.get(w, 0) / top for w in ws) / (len(ws) ** 0.6)
        if i < 3:
            score *= 1.15
        if len(ws) > 60:
            score *= 0.7
        scored.append((score, i, s))
    best = sorted(sorted(scored, reverse=True)[:k], key=lambda x: x[1])
    summary = [s for _, _, s in best]
    keywords = [w for w, _ in freq.most_common(12)]
    return {"summary": " ".join(summary), "points": summary[:min(8, len(summary))], "keywords": keywords,
            "stats": {"words": len(re.findall(r"\w+", text)), "sentences": len(sents),
                      "summary_words": len(re.findall(r"\w+", " ".join(summary)))}}


def summarize(files, opts, wd):
    up = require_pdfs(files, 1, 1)[0]
    doc = open_pdf(up)
    text = "\n".join(p.get_text() for p in doc)
    res = summarize_text(text, opts.get("length", "medium"))
    txt = (f"Summary of {up.name}\n\n{res['summary']}\n\nKey points:\n" +
           "\n".join(f"• {p}" for p in res["points"]) + "\n\nKeywords: " + ", ".join(res["keywords"]))
    res["download"] = base64.b64encode(txt.encode()).decode()
    res["download_name"] = f"{up.stem}_summary.txt"
    return Result(data=res)


LANGS = {"af": "Afrikaans", "ar": "Arabic", "bn": "Bengali", "zh-CN": "Chinese (Simplified)", "zh-TW": "Chinese (Traditional)",
         "cs": "Czech", "da": "Danish", "nl": "Dutch", "en": "English", "fi": "Finnish", "fr": "French", "de": "German",
         "el": "Greek", "gu": "Gujarati", "he": "Hebrew", "hi": "Hindi", "hu": "Hungarian", "id": "Indonesian",
         "it": "Italian", "ja": "Japanese", "ko": "Korean", "ms": "Malay", "mr": "Marathi", "ne": "Nepali", "no": "Norwegian",
         "fa": "Persian", "pl": "Polish", "pt": "Portuguese", "pa": "Punjabi", "ro": "Romanian", "ru": "Russian",
         "es": "Spanish", "sv": "Swedish", "ta": "Tamil", "te": "Telugu", "th": "Thai", "tr": "Turkish",
         "uk": "Ukrainian", "ur": "Urdu", "vi": "Vietnamese"}


def _translate_one(text, target, source="auto"):
    text = text.strip()
    if not text or not re.search(r"[^\W\d_]", text):
        return text
    last_err = None
    for attempt in range(3):
        try:
            data = urllib.parse.urlencode({"q": text}).encode()
            url = ("https://translate.googleapis.com/translate_a/single?client=gtx&dt=t&"
                   + urllib.parse.urlencode({"sl": source, "tl": target}))
            req = urllib.request.Request(url, data=data, headers={"User-Agent": "Mozilla/5.0",
                                                                   "Content-Type": "application/x-www-form-urlencoded;charset=utf-8"})
            res = json.loads(urllib.request.urlopen(req, timeout=20).read().decode("utf-8"))
            return "".join(seg[0] for seg in res[0] if seg and seg[0])
        except Exception as e:
            last_err = e
    try:  # fallback service
        url = "https://api.mymemory.translated.net/get?" + urllib.parse.urlencode(
            {"q": text[:480], "langpair": f"{'en' if source == 'auto' else source}|{target}"})
        res = json.loads(urllib.request.urlopen(url, timeout=20).read().decode("utf-8"))
        t = res.get("responseData", {}).get("translatedText")
        if t:
            return t
    except Exception as e:
        last_err = e
    raise UserError(f"Translation service unreachable (internet connection required): {last_err}")


def translate(files, opts, wd):
    up = require_pdfs(files, 1, 1)[0]
    target = opts.get("target") or "bn"
    if target not in LANGS:
        raise UserError("Unsupported target language.")
    source = opts.get("source") or "auto"
    doc = open_pdf(up)
    jobs = []
    for page in doc:
        normalize_page(page)
        for b in page.get_text("dict", flags=pymupdf.TEXT_PRESERVE_WHITESPACE)["blocks"]:
            if b["type"] != 0:
                continue
            spans = [s for l in b["lines"] for s in l["spans"] if s["text"].strip()]
            if not spans:
                continue
            text = " ".join("".join(s["text"] for s in l["spans"]).strip() for l in b["lines"]).strip()
            text = re.sub(r"(\w)- (\w)", r"\1\2", text)
            main = max(spans, key=lambda s: len(s["text"]))
            color = "#%06x" % main["color"]
            bold = bool(main["flags"] & 16) or "bold" in main["font"].lower()
            serif = bool(main["flags"] & 4)
            jobs.append({"page": page.number, "bbox": pymupdf.Rect(b["bbox"]), "text": text,
                         "size": main["size"], "color": color, "bold": bold, "serif": serif})
    if not jobs:
        raise UserError("No text found to translate. If this is a scanned PDF, run OCR PDF first.")
    if sum(len(j["text"]) for j in jobs) > 300000:
        raise UserError("This document is too long to translate at once (limit ~300,000 characters).")
    cache = {}
    uniq = list({j["text"] for j in jobs})
    with cf.ThreadPoolExecutor(max_workers=8) as ex:
        for t, res in zip(uniq, ex.map(lambda s: _translate_one(s, target, source), uniq)):
            cache[t] = res
    by_page = {}
    for j in jobs:
        by_page.setdefault(j["page"], []).append(j)
    for pno, lst in by_page.items():
        page = doc[pno]
        for j in lst:
            page.add_redact_annot(j["bbox"], fill=False, cross_out=False)
        page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE, graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
                              text=pymupdf.PDF_REDACT_TEXT_REMOVE)
        for j in lst:
            rect = pymupdf.Rect(j["bbox"])
            rect.x1 = min(page.rect.x1, rect.x1 + 2)
            rect.y1 = min(page.rect.y1, rect.y1 + j["size"] * 0.6)
            insert_text_box(page, rect, cache[j["text"]], size=j["size"], color=j["color"],
                            family="serif" if j["serif"] else "sans", bold=j["bold"])
    return Result([(f"{up.stem}_{target}.pdf", pdf_bytes(doc))], {"blocks": len(jobs)})


# =========================================================================== #
# Registry
# =========================================================================== #
TOOLS = {
    "merge": merge, "split": split, "remove-pages": remove_pages, "extract-pages": extract_pages,
    "organize": organize, "rotate": rotate, "compress": compress, "repair": repair, "ocr": ocr,
    "jpg-to-pdf": images_to_pdf, "scan-to-pdf": images_to_pdf,
    "word-to-pdf": office_to_pdf("word"), "powerpoint-to-pdf": office_to_pdf("powerpoint"),
    "excel-to-pdf": office_to_pdf("excel"), "html-to-pdf": html_to_pdf,
    "pdf-to-jpg": pdf_to_images, "pdf-to-word": pdf_to_word, "pdf-to-powerpoint": pdf_to_powerpoint,
    "pdf-to-excel": pdf_to_excel, "pdf-to-pdfa": pdf_to_pdfa, "pdf-to-markdown": pdf_to_markdown,
    "page-numbers": page_numbers, "watermark": watermark, "crop": crop, "edit": edit, "sign": edit,
    "forms": forms, "form-fields": form_fields, "unlock": unlock, "protect": protect, "redact": redact,
    "compare": compare, "summarize": summarize, "translate": translate,
}

# Tools allowed inside a workflow (single PDF in -> single PDF out)
WORKFLOW_TOOLS = {"compress", "rotate", "watermark", "page-numbers", "protect", "unlock", "repair", "ocr",
                  "pdf-to-pdfa", "crop", "redact", "remove-pages", "extract-pages", "translate"}


def workflow(files, opts, wd):
    steps = opts.get("steps") or []
    if not steps:
        raise UserError("This workflow has no steps.")
    for s in steps:
        if s.get("tool") not in WORKFLOW_TOOLS:
            raise UserError(f"Tool “{s.get('tool')}” can't be used in a workflow.")
    if steps[-1:] and any(s["tool"] == "protect" for s in steps[:-1]):
        raise UserError("“Protect PDF” must be the last step of a workflow.")
    pdfs = require_pdfs(files, 1)
    out = []
    for up in pdfs:
        cur = up
        for i, s in enumerate(steps):
            res = TOOLS[s["tool"]]([cur], s.get("options") or {}, wd)
            if len(res.files) != 1:
                raise UserError(f"Step {i + 1} ({s['tool']}) produced {len(res.files)} files – workflows need one PDF per step.")
            p = os.path.join(wd, f"wf_{secrets.token_hex(4)}.pdf")
            Path(p).write_bytes(res.files[0][1])
            cur = Upload(up.name, p)
        out.append((f"{up.stem}_workflow.pdf", Path(cur.path).read_bytes()))
    return Result(out)


TOOLS["workflow"] = workflow
